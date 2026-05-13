#!/usr/bin/env python3
"""
Generate: Assignment 13 â€” Country Profile: Japan
Professional .docx for college submission
Author: William Robert Schreiner Jr
"""

import os

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# â”€â”€ COLOR PALETTE â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
DARK_RED  = RGBColor(0x8B, 0x00, 0x00)
CHARCOAL  = RGBColor(0x2E, 0x2E, 0x2E)
MID_GRAY  = RGBColor(0x75, 0x75, 0x75)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)

FONT      = "Calibri"
OUT_PATH  = r"c:\Users\willf\OneDrive\Documents\Python Class\word\Assignment13_CountryProfile_Japan_Final.docx"



# â”€â”€ IMAGE DOWNLOADER â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€



# â”€â”€ XML / TABLE HELPERS â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

def cell_shading(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for ex in tcPr.findall(qn("w:shd")):
        tcPr.remove(ex)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, sides: dict):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for ex in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(ex)
    tcBd = OxmlElement("w:tcBorders")
    for side, (val, sz, color) in sides.items():
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"),   val)
        el.set(qn("w:sz"),    sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBd.append(el)
    tcPr.append(tcBd)


NO_BORDER = {s: ("none", "0", "auto") for s in ("top", "bottom", "left", "right")}


def set_cell_padding(cell, top=80, bottom=80, left=160, right=80):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(val));  m.set(qn("w:type"), "dxa")
        tcMar.append(m)
    tcPr.append(tcMar)


def set_cell_valign(cell, align="center"):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    va   = OxmlElement("w:vAlign")
    va.set(qn("w:val"), align)
    tcPr.append(va)


def hr_rule(doc, color="8B0000", sz="6"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    sz)
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), color)
    pBdr.append(bot)
    pPr.append(pBdr)
    return p


def set_run_fmt(run, size=8.5):
    run.font.name  = FONT
    run.font.size  = Pt(size)
    run.font.color.rgb = MID_GRAY


# â”€â”€ LAYOUT HELPERS â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

def section_banner(doc, num: str, title: str):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    c0 = tbl.cell(0, 0);  c0.width = Inches(0.55)
    cell_shading(c0, "8B0000");  set_cell_borders(c0, NO_BORDER)
    set_cell_padding(c0, 120, 120, 120, 80);  set_cell_valign(c0, "center")
    p0 = c0.paragraphs[0];  p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run(num)
    r0.font.name = FONT;  r0.font.size = Pt(18);  r0.font.bold = True;  r0.font.color.rgb = WHITE
    c1 = tbl.cell(0, 1);  c1.width = Inches(5.45)
    cell_shading(c1, "2E2E2E");  set_cell_borders(c1, NO_BORDER)
    set_cell_padding(c1, 120, 120, 200, 120);  set_cell_valign(c1, "center")
    p1 = c1.paragraphs[0]
    r1 = p1.add_run(title.upper())
    r1.font.name = FONT;  r1.font.size = Pt(13);  r1.font.bold = True;  r1.font.color.rgb = WHITE
    sp = doc.add_paragraph();  sp.paragraph_format.space_before = Pt(12);  sp.paragraph_format.space_after = Pt(0)


def body_p(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(11)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.first_line_indent = Inches(0.25)
    r = p.add_run(text)
    r.font.name = FONT;  r.font.size = Pt(11.5);  r.font.color.rgb = CHARCOAL
    return p


def quote_box(doc, text: str):
    tbl = doc.add_table(rows=1, cols=1);  tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = tbl.cell(0, 0)
    cell_shading(cell, "FBF0F0")
    set_cell_borders(cell, {
        "top": ("none","0","auto"), "bottom": ("none","0","auto"),
        "right": ("none","0","auto"), "left": ("single","28","8B0000"),
    })
    set_cell_padding(cell, 120, 120, 220, 140)
    p = cell.paragraphs[0]
    rq = p.add_run(text)
    rq.font.name = FONT;  rq.font.size = Pt(11.5);  rq.font.italic = True;  rq.font.color.rgb = DARK_RED
    sp = doc.add_paragraph();  sp.paragraph_format.space_before = Pt(14);  sp.paragraph_format.space_after = Pt(0)


def captioned_image(doc, img_path: str, caption: str, width=Inches(5.5)):
    """Centered image with italic gray caption below."""
    pi = doc.add_paragraph()
    pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pi.paragraph_format.space_before = Pt(10)
    pi.paragraph_format.space_after  = Pt(2)
    pi.add_run().add_picture(img_path, width=width)
    pc = doc.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.paragraph_format.space_before = Pt(0)
    pc.paragraph_format.space_after  = Pt(14)
    rc = pc.add_run(caption)
    rc.font.name = FONT;  rc.font.size = Pt(9.5)
    rc.font.italic = True;  rc.font.color.rgb = MID_GRAY


def side_by_side_text_image(doc, text: str, img_path: str, caption: str,
                             text_width=Inches(3.4), img_width_col=Inches(2.6),
                             img_render_width=Inches(2.35)):
    """2-column table: body text on left, image + caption on right."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    # â”€â”€ text cell â”€â”€
    ct = tbl.cell(0, 0);  ct.width = text_width
    cell_shading(ct, "FFFFFF");  set_cell_borders(ct, NO_BORDER)
    set_cell_padding(ct, 40, 40, 0, 120);  set_cell_valign(ct, "top")
    pt = ct.paragraphs[0]
    pt.paragraph_format.space_before = Pt(0)
    pt.paragraph_format.space_after  = Pt(8)
    pt.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    rt = pt.add_run(text)
    rt.font.name = FONT;  rt.font.size = Pt(11);  rt.font.color.rgb = CHARCOAL

    # â”€â”€ image cell â”€â”€
    ci = tbl.cell(0, 1);  ci.width = img_width_col
    cell_shading(ci, "FFFFFF");  set_cell_borders(ci, NO_BORDER)
    set_cell_padding(ci, 40, 40, 80, 0);  set_cell_valign(ci, "top")
    pi_p = ci.paragraphs[0]
    pi_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pi_p.paragraph_format.space_before = Pt(0)
    pi_p.paragraph_format.space_after  = Pt(4)
    pi_p.add_run().add_picture(img_path, width=img_render_width)
    # caption paragraph
    pc = ci.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.paragraph_format.space_before = Pt(2)
    pc.paragraph_format.space_after  = Pt(0)
    rc = pc.add_run(caption)
    rc.font.name = FONT;  rc.font.size = Pt(8.5)
    rc.font.italic = True;  rc.font.color.rgb = MID_GRAY

    sp = doc.add_paragraph();  sp.paragraph_format.space_before = Pt(10);  sp.paragraph_format.space_after = Pt(0)


def build_footer(doc):
    sec    = doc.sections[0]
    footer = sec.footer
    fp     = footer.paragraphs[0]
    fp.clear()
    tbl = footer.add_table(1, 2, Inches(6.0))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in tbl.rows:
        for cell in row.cells:
            set_cell_borders(cell, NO_BORDER)
    cl = tbl.cell(0, 0);  cl.width = Inches(3.5)
    pl = cl.paragraphs[0];  pl.paragraph_format.space_before = Pt(0)
    rl = pl.add_run("William Robert Schreiner Jr  \u2022  Assignment 13");  set_run_fmt(rl)
    cr = tbl.cell(0, 1);  cr.width = Inches(2.5)
    pr = cr.paragraphs[0];  pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT;  pr.paragraph_format.space_before = Pt(0)
    rr0 = pr.add_run("Page ");  set_run_fmt(rr0)
    for fld_type in ("begin", "PAGE", "end"):
        r = pr.add_run();  set_run_fmt(r)
        if fld_type in ("begin", "end"):
            fc = OxmlElement("w:fldChar");  fc.set(qn("w:fldCharType"), fld_type);  r._r.append(fc)
        else:
            instr = OxmlElement("w:instrText");  instr.set(qn("xml:space"), "preserve");  instr.text = " PAGE ";  r._r.append(instr)


# â”€â”€ MAIN DOCUMENT BUILD â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€

def build():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin    = Inches(1.0)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin   = Inches(1.1)
        sec.right_margin  = Inches(1.1)

    build_footer(doc)


    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    # TITLE PAGE
    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•

    # Dark red top banner
    tbl_top = doc.add_table(rows=1, cols=1);  tbl_top.alignment = WD_TABLE_ALIGNMENT.LEFT
    ctop = tbl_top.cell(0, 0);  ctop.width = Inches(6.0)
    cell_shading(ctop, "8B0000");  set_cell_borders(ctop, NO_BORDER)
    set_cell_padding(ctop, 200, 200, 200, 200)
    ptop = ctop.paragraphs[0];  ptop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rtop = ptop.add_run("MACROECONOMICS  \u2022  COUNTRY PROFILE ASSIGNMENT")
    rtop.font.name = FONT;  rtop.font.size = Pt(10);  rtop.font.bold = True;  rtop.font.color.rgb = WHITE
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # Spacers
    for _ in range(2):
        sp = doc.add_paragraph();  sp.paragraph_format.space_before = Pt(8);  sp.paragraph_format.space_after = Pt(0)

    # Title block
    pt1 = doc.add_paragraph();  pt1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt1 = pt1.add_run("Assignment 13")
    rt1.font.name = FONT;  rt1.font.size = Pt(42);  rt1.font.bold = True;  rt1.font.color.rgb = DARK_RED
    pt2 = doc.add_paragraph();  pt2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rt2 = pt2.add_run("Country Profile \u2013 Japan")
    rt2.font.name = FONT;  rt2.font.size = Pt(24);  rt2.font.color.rgb = CHARCOAL

    hr_rule(doc, "8B0000", "10")


    hr_rule(doc, "C0C0C0", "4")

    sp_info = doc.add_paragraph();  sp_info.paragraph_format.space_before = Pt(18);  sp_info.paragraph_format.space_after = Pt(0)

    # Prepared by / date block
    info_tbl = doc.add_table(rows=2, cols=2);  info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ["Prepared by", "Date Submitted"]
    values = ["William Robert Schreiner Jr", "May 2026"]
    for i in range(2):
        lc = info_tbl.cell(i, 0);  vc = info_tbl.cell(i, 1)
        for c in (lc, vc):
            set_cell_borders(c, NO_BORDER);  set_cell_padding(c, 40, 40, 60, 60)
        lc.width = Inches(1.6);  vc.width = Inches(3.0)
        pl = lc.paragraphs[0];  rl = pl.add_run(labels[i] + ":")
        rl.font.name = FONT;  rl.font.size = Pt(10.5);  rl.font.bold = True;  rl.font.color.rgb = MID_GRAY
        pv = vc.paragraphs[0];  rv = pv.add_run(values[i])
        rv.font.name = FONT;  rv.font.size = Pt(10.5);  rv.font.color.rgb = CHARCOAL

    sp_b = doc.add_paragraph();  sp_b.paragraph_format.space_before = Pt(20)
    hr_rule(doc, "C0C0C0", "4")
    doc.add_page_break()

    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    # SECTION 1 â€” COUNTRY OVERVIEW
    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    section_banner(doc, "01", "Country Overview")

    pcl = doc.add_paragraph()
    pcl.paragraph_format.space_before = Pt(0);  pcl.paragraph_format.space_after = Pt(9)
    rcl1 = pcl.add_run("Country Chosen: ")
    rcl1.font.name = FONT;  rcl1.font.size = Pt(11.5);  rcl1.font.bold = True;  rcl1.font.color.rgb = CHARCOAL
    rcl2 = pcl.add_run("Japan")
    rcl2.font.name = FONT;  rcl2.font.size = Pt(11.5);  rcl2.font.bold = True;  rcl2.font.color.rgb = DARK_RED

    overview_text = (
        "Japan was selected for this country profile because it stands out as one of the most "
        "technologically advanced nations in the world. The country has made significant "
        "contributions to global innovation, particularly in the areas of robotics, consumer "
        "electronics, and transportation \u2014 most notably through the iconic Shinkansen bullet "
        "train system, internationally recognized for its speed, precision, and reliability. "
        "What makes Japan especially compelling is its remarkable ability to blend cutting-edge "
        "technology with deeply rooted cultural traditions that have persisted for centuries. "
        "Japan is also home to some of the world\u2019s most influential corporations, including "
        "Toyota, Sony, and Nintendo \u2014 companies that have reshaped entire industries globally. "
        "Additionally, Japan consistently ranks among countries with the highest life expectancy "
        "in the world, reflecting its strong public health infrastructure and high quality of life."
    )

    body_p(doc, overview_text)

    quote_box(doc,
        "\u201c Japan\u2019s global leadership in robotics, electronics, and high-speed "
        "transportation makes it one of the most innovative and technologically sophisticated "
        "economies in the modern world. \u201d"
    )

    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    # SECTION 2 â€” ECONOMIC ANALYSIS
    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    section_banner(doc, "02", "Economic Analysis")

    body_p(doc,
        "A recession is generally defined as a period of significant economic decline, typically "
        "characterized by falling production, reduced consumer spending, and rising unemployment. "
        "Among economists, the most widely accepted measure of a recession is two consecutive "
        "quarters of negative GDP (Gross Domestic Product) growth. While recessions can vary in "
        "length and severity, they signal a broader contraction in economic activity that affects "
        "businesses, workers, and households alike."
    )
    body_p(doc,
        "As of the time this profile was prepared, Japan is not considered to be in a recession. "
        "Recent economic data has pointed to positive overall growth, suggesting the economy "
        "remains in an expansion phase. That said, Japan continues to face notable structural "
        "challenges, including persistent inflationary pressures and a rapidly aging population "
        "\u2014 both of which present long-term concerns for economic sustainability and labor "
        "supply. Despite these headwinds, Japan\u2019s economy has continued to grow, supported "
        "by strong export industries, government investment, and a resilient domestic market."
    )

    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    # SECTION 3 â€” KEY FACTS ABOUT JAPAN
    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    section_banner(doc, "03", "Key Facts About Japan")

    pi2 = doc.add_paragraph()
    pi2.paragraph_format.space_before = Pt(0);  pi2.paragraph_format.space_after = Pt(10)
    ri2 = pi2.add_run("The table below summarizes key demographic, economic, and cultural facts about Japan.")
    ri2.font.name = FONT;  ri2.font.size = Pt(11.5);  ri2.font.italic = True;  ri2.font.color.rgb = CHARCOAL

    facts = [
        ("Capital",               "Tokyo"),
        ("Population",            "Approximately 124 million"),
        ("Currency",              "Japanese Yen (\u00a5)"),
        ("Major Industries",      "Automotive \u2022 Electronics \u2022 Robotics \u2022 Manufacturing"),
        ("Major Companies",       "Toyota \u2022 Sony \u2022 Nintendo"),
        ("Famous Transportation", "Shinkansen Bullet Train \u2014 world-renowned for speed & reliability"),
        ("Life Expectancy",       "Among the highest in the world"),
    ]

    tbl = doc.add_table(rows=len(facts) + 1, cols=2)
    tbl.style = "Table Grid";  tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    COL0_W = Inches(1.95);  COL1_W = Inches(4.05)

    hdr = tbl.rows[0]
    for j, htxt in enumerate(["Category", "Details"]):
        hc = hdr.cells[j];  hc.width = COL0_W if j == 0 else COL1_W
        cell_shading(hc, "8B0000");  set_cell_padding(hc, 140, 140, 160, 100);  set_cell_valign(hc, "center")
        hr2 = hc.paragraphs[0].add_run(htxt)
        hr2.font.name = FONT;  hr2.font.size = Pt(11.5);  hr2.font.bold = True;  hr2.font.color.rgb = WHITE

    for i, (cat, det) in enumerate(facts):
        row = tbl.rows[i + 1];  c0, c1 = row.cells[0], row.cells[1]
        c0.width = COL0_W;  c1.width = COL1_W
        cell_shading(c0, "F0EDED");  cell_shading(c1, "F5F5F5" if i % 2 == 0 else "FFFFFF")
        set_cell_padding(c0, 120, 120, 160, 100);  set_cell_padding(c1, 120, 120, 160, 100)
        set_cell_valign(c0, "center");  set_cell_valign(c1, "center")
        r0 = c0.paragraphs[0].add_run(cat)
        r0.font.name = FONT;  r0.font.size = Pt(11);  r0.font.bold = True;  r0.font.color.rgb = DARK_RED
        r1 = c1.paragraphs[0].add_run(det)
        r1.font.name = FONT;  r1.font.size = Pt(11);  r1.font.color.rgb = CHARCOAL

    # Spacer after table
    sp = doc.add_paragraph();  sp.paragraph_format.space_before = Pt(16)

    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    # SECTION 4 â€” REFERENCES
    # â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•â•
    section_banner(doc, "04", "References")

    sources = [
        ("International Monetary Fund. (2024).",
         "World Economic Outlook Database.",
         "International Monetary Fund. https://www.imf.org/en/Publications/WEO"),
        ("The World Bank. (2024).",
         "World Bank Open Data \u2013 Japan.",
         "The World Bank Group. https://data.worldbank.org/country/japan"),
        ("BBC News. (2024).",
         "Japan economy: Business reports and analysis.",
         "BBC News. https://www.bbc.com/news/business"),
    ]

    for author, title_txt, rest in sources:
        ps = doc.add_paragraph()
        ps.paragraph_format.space_before = Pt(0);  ps.paragraph_format.space_after = Pt(9)
        ps.paragraph_format.left_indent = Inches(0.5);  ps.paragraph_format.first_line_indent = Inches(-0.5)
        ra = ps.add_run(author + " ");  ra.font.name = FONT;  ra.font.size = Pt(11);  ra.font.color.rgb = CHARCOAL
        rt_src = ps.add_run(title_txt + " ");  rt_src.font.name = FONT;  rt_src.font.size = Pt(11);  rt_src.font.italic = True;  rt_src.font.color.rgb = CHARCOAL
        rr = ps.add_run(rest);  rr.font.name = FONT;  rr.font.size = Pt(11);  rr.font.color.rgb = CHARCOAL

    # â”€â”€ SAVE â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€â”€
    doc.save(OUT_PATH)
    print(f"\u2714  Saved: {OUT_PATH}\n")



if __name__ == "__main__":
    build()
